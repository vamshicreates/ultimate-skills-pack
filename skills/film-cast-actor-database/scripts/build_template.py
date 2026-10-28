"""Regenerate the blank workbook shipped with the skill."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "templates" / "actor_database_template.xlsx"

SHEETS = [
    (
        "Actors",
        [
            "Actor name",
            "Potential possible roles",
            "Real age",
            "Checked on date",
            "Instagram link",
            "Source link",
        ],
        [28, 42, 14, 18, 42, 52],
    ),
    (
        "Film evidence",
        [
            "Actor name",
            "Recent film #",
            "Film",
            "Release year",
            "Credited character",
            "Credit note",
            "Film source",
        ],
        [28, 17, 43, 17, 33, 35, 52],
    ),
]


def main() -> None:
    wb = Workbook()
    for index, (title, headers, widths) in enumerate(SHEETS):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = title
        ws.freeze_panes = "B2"
        ws.append(headers)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1000"
        ws.row_dimensions[1].height = 27
        for column, (header, width) in enumerate(zip(headers, widths), start=1):
            cell = ws.cell(row=1, column=column)
            cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="356855")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(column)].width = width
        ws.sheet_view.showGridLines = False
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
